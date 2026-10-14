"""
Jira Server / Data Center worklog çekirdeği (arayüzden bağımsız).

- Jira'ya yalnızca GET (okuma) isteği atar. Yazma/silme yapmaz.
- Token yalnızca bellekte tutulur; hiçbir dosyaya yazılmaz.
- Sadece Python standart kütüphanesini kullanır (xlsx çıktısı için openpyxl hariç).
"""
from __future__ import annotations

import csv
import json
import re
import socket
import ssl
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timedelta, timezone
from pathlib import Path
from typing import Callable, Optional

DEFAULT_URL = ""  # kasıtlı olarak boş: adres ilk kullanımda sorulur ve yerel ayar dosyasına kaydedilir

COLUMNS = ["Kullanıcı", "Tarih", "Başlama Saati", "Task", "Özet", "Proje", "Tip", "Saat", "Yorum", "Girilme Zamanı"]

# Türkiye 2016'dan beri yıl boyunca UTC+3 (yaz/kış saati uygulaması yok); ek bileşen gerekmez.
TURKEY_TZ = timezone(timedelta(hours=3))
_JIRA_DT_RE = re.compile(r"^(\d{4}-\d{2}-\d{2}T\d{2}:\d{2}:\d{2})(?:\.\d+)?(Z|[+-]\d{2}:?\d{2})?$")

_PROJECT_RE = re.compile(r"^[A-Za-z][A-Za-z0-9_]*$")
_URL_CUT_MARKERS = ("/browse", "/secure", "/projects", "/issues", "/rest", "/plugins", "/login", "/dashboard")

ProgressFn = Callable[[str, int, int], None]  # (aşama, tamamlanan, toplam)


class JiraError(Exception):
    """Kullanıcıya olduğu gibi gösterilebilecek, anlaşılır hata."""


class Cancelled(Exception):
    """Kullanıcı işlemi iptal etti."""


# --------------------------------------------------------------------------
# Girdi doğrulama / düzeltme
# --------------------------------------------------------------------------
def parse_date(text: str) -> date:
    text = (text or "").strip()
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise JiraError(f"Tarih anlaşılamadı: '{text}'. Örnek: 01.09.2026")


def parse_projects(text: str) -> list[str]:
    keys = [k.strip().upper() for k in re.split(r"[,\s;]+", text or "") if k.strip()]
    if not keys:
        raise JiraError("Proje anahtarı boş olamaz. Örnek: PROJ (birden fazlası için virgülle ayır).")
    for k in keys:
        if not _PROJECT_RE.match(k):
            raise JiraError(
                f"Proje anahtarı geçersiz: '{k}'. Sadece harf, rakam ve alt çizgi olmalı (örnek: PROJ)."
            )
    return keys


def check_range(start: date, end: date) -> None:
    if end < start:
        raise JiraError("Bitiş tarihi başlangıç tarihinden önce olamaz.")


def normalize_url(text: str) -> str:
    """'https://jira.sirket.com/browse/ABC-1?filter=-4' gibi yapıştırılan adresi ana adrese indirir."""
    text = (text or "").strip()
    if not text:
        raise JiraError("Jira adresi boş olamaz.")
    if not re.match(r"^https?://", text, re.I):
        text = "https://" + text
    parsed = urllib.parse.urlparse(text)
    if not parsed.netloc:
        raise JiraError(f"Jira adresi anlaşılamadı: '{text}'.")
    path = parsed.path
    for marker in _URL_CUT_MARKERS:
        idx = path.find(marker)
        if idx != -1:
            path = path[:idx]
    return f"{parsed.scheme}://{parsed.netloc}{path}".rstrip("/")


def clean_token(token: str) -> str:
    token = (token or "").strip()
    if token.lower().startswith("bearer "):
        token = token[7:].strip()
    if not token:
        raise JiraError("Token boş olamaz.")
    return token


def previous_month_range(today: Optional[date] = None) -> tuple[date, date]:
    today = today or date.today()
    last_prev = today.replace(day=1) - timedelta(days=1)
    return last_prev.replace(day=1), last_prev


def current_month_range(today: Optional[date] = None) -> tuple[date, date]:
    today = today or date.today()
    return today.replace(day=1), today


def parse_jira_datetime(text: str) -> datetime:
    """'2026-09-03T10:15:00.000+0300' -> saat dilimi bilgili datetime. Ofset yoksa Türkiye saati varsayılır."""
    m = _JIRA_DT_RE.match((text or "").strip())
    if not m:
        raise JiraError(f"Jira'dan beklenmeyen bir tarih/saat geldi: '{text}'.")
    naive = datetime.strptime(m.group(1), "%Y-%m-%dT%H:%M:%S")
    off = m.group(2)
    if off is None:
        tz = TURKEY_TZ
    elif off == "Z":
        tz = timezone.utc
    else:
        digits = off[1:].replace(":", "")
        delta = timedelta(hours=int(digits[:2]), minutes=int(digits[2:]))
        tz = timezone(-delta if off[0] == "-" else delta)
    return naive.replace(tzinfo=tz)


def to_turkey_time(text: str) -> datetime:
    """Jira'nın verdiği zamanı Türkiye saatine çevirir (saat dilimsiz, duvar saati olarak döner)."""
    return parse_jira_datetime(text).astimezone(TURKEY_TZ).replace(tzinfo=None)


# --------------------------------------------------------------------------
# Jira istemcisi (yalnızca GET)
# --------------------------------------------------------------------------
class _SafeRedirect(urllib.request.HTTPRedirectHandler):
    """Token'ın başka bir sunucuya (ya da şifresiz http'ye) yönlendirilmesini engeller."""

    def redirect_request(self, req, fp, code, msg, headers, newurl):
        old = urllib.parse.urlparse(req.full_url)
        new = urllib.parse.urlparse(newurl)
        if new.netloc != old.netloc or (old.scheme == "https" and new.scheme != "https"):
            raise JiraError(
                "Jira başka bir adrese yönlendirdi; güvenlik için durduruldu. "
                "Jira adresini tarayıcıdaki ana adresle aynı yazdığından emin ol."
            )
        return super().redirect_request(req, fp, code, msg, headers, newurl)


class JiraClient:
    def __init__(self, base_url: str, token: str, timeout: int = 60):
        self.base_url = normalize_url(base_url)
        self._token = clean_token(token)
        self.timeout = timeout
        self._openers = self._build_openers()
        self._opener_idx = 0

    @staticmethod
    def _build_openers() -> list:
        contexts = [ssl.create_default_context()]  # önce sistemin sertifika deposu
        try:
            import certifi  # paketlenmiş uygulamada sistem deposu olmayabilir

            contexts.append(ssl.create_default_context(cafile=certifi.where()))
        except Exception:
            pass
        return [
            urllib.request.build_opener(urllib.request.HTTPSHandler(context=c), _SafeRedirect())
            for c in contexts
        ]

    # -- düşük seviye -------------------------------------------------------
    def get_json(self, path: str, params: Optional[dict] = None):
        url = self.base_url + path
        if params:
            url += "?" + urllib.parse.urlencode(params)
        attempt = 0
        while True:
            req = urllib.request.Request(
                url,
                method="GET",  # bu uygulamada başka bir HTTP metodu kullanılmaz
                headers={"Authorization": "Bearer " + self._token, "Accept": "application/json"},
            )
            try:
                with self._openers[self._opener_idx].open(req, timeout=self.timeout) as resp:
                    return json.load(resp)
            except urllib.error.HTTPError as e:
                if e.code in (429, 502, 503, 504) and attempt < 3:
                    try:
                        delay = float(e.headers.get("Retry-After", ""))
                    except ValueError:
                        delay = 2**attempt
                    time.sleep(min(delay, 30))
                    attempt += 1
                    continue
                raise self._http_error(e) from None
            except urllib.error.URLError as e:
                if isinstance(e.reason, ssl.SSLCertVerificationError):
                    if self._opener_idx + 1 < len(self._openers):
                        self._opener_idx += 1
                        continue
                    raise JiraError(
                        "Güvenlik sertifikası doğrulanamadı. Şirket ağında/VPN'de olduğundan "
                        "emin ol; sorun sürerse BT ile görüş."
                    ) from None
                raise JiraError(
                    "Jira'ya ulaşılamadı. İnternet/VPN bağlantını ve Jira adresini kontrol et."
                ) from None
            except (TimeoutError, socket.timeout):
                raise JiraError("Jira zamanında yanıt vermedi. Biraz sonra tekrar dene.") from None
            except json.JSONDecodeError:
                raise JiraError(
                    "Jira beklenmeyen bir yanıt verdi (giriş sayfası olabilir). "
                    "Jira adresini ve token'ı kontrol et."
                ) from None

    @staticmethod
    def _http_error(e: urllib.error.HTTPError) -> JiraError:
        detail = ""
        try:
            body = json.loads(e.read().decode("utf-8", "replace"))
            msgs = body.get("errorMessages") or []
            detail = " ".join(msgs) if msgs else ""
        except Exception:
            pass
        if e.code == 401:
            return JiraError("Token geçersiz, süresi dolmuş ya da yanlış yapıştırılmış (401).")
        if e.code == 403:
            return JiraError("Bu işlem için yetkin yok (403). Projeye erişimin olduğundan emin ol.")
        if e.code == 404:
            return JiraError("Adres bulunamadı (404). Jira adresini kontrol et.")
        if e.code == 400:
            return JiraError("Jira sorguyu kabul etmedi" + (f": {detail}" if detail else " (400)."))
        return JiraError(f"Jira hata verdi (HTTP {e.code})." + (f" {detail}" if detail else ""))

    # -- yüksek seviye ------------------------------------------------------
    def search_issues(self, projects: list[str], start: date, end: date, cancel=None) -> list[dict]:
        proj = ", ".join(f'"{p}"' for p in projects)
        jql = f'project in ({proj}) AND worklogDate >= "{start:%Y-%m-%d}" AND worklogDate <= "{end:%Y-%m-%d}"'
        issues: list[dict] = []
        start_at = 0
        while True:
            _check_cancel(cancel)
            data = self.get_json(
                "/rest/api/2/search",
                {"jql": jql, "fields": "summary,project,issuetype", "startAt": start_at, "maxResults": 100},
            )
            batch = data.get("issues", [])
            issues.extend(batch)
            start_at += len(batch)
            if not batch or start_at >= data.get("total", 0):
                break
        return issues

    def issue_worklogs(self, key: str) -> list[dict]:
        path = f"/rest/api/2/issue/{urllib.parse.quote(key)}/worklog"
        data = self.get_json(path)
        out = list(data.get("worklogs", []))
        # Çok nadir: kayıtlar tek yanıta sığmadıysa kalanını sayfalayarak al
        while out and len(out) < data.get("total", len(out)):
            more = self.get_json(path, {"startAt": len(out), "maxResults": 1000}).get("worklogs", [])
            if not more:
                break
            out.extend(more)
        return out


def _check_cancel(cancel) -> None:
    if cancel is not None and cancel.is_set():
        raise Cancelled()


def fetch_report(
    base_url: str,
    token: str,
    projects: list[str],
    start: date,
    end: date,
    progress: Optional[ProgressFn] = None,
    cancel=None,
) -> list[dict]:
    """Seçilen proje(ler) ve tarih aralığı için worklog satırlarını döndürür."""
    check_range(start, end)
    client = JiraClient(base_url, token)
    say = progress or (lambda stage, done, total: None)

    say("search", 0, 0)
    # Jira "worklogDate"i kendi saat diliminde değerlendirir; Türkiye saatiyle gün sınırında kalan kayıtlar
    # kaçmasın diye aralığı bir gün genişletip kesin filtreyi aşağıda Türkiye saatine göre uygularız.
    issues = client.search_issues(projects, start - timedelta(days=1), end + timedelta(days=1), cancel)
    total = len(issues)
    say("issues", 0, total)

    rows: list[dict] = []
    for n, issue in enumerate(issues, 1):
        _check_cancel(cancel)
        key = issue["key"]
        fields = issue.get("fields", {})
        for w in client.issue_worklogs(key):
            started = to_turkey_time(w["started"])  # Jira'nın verdiği ofset okunur, Türkiye saatine çevrilir
            if not (start <= started.date() <= end):
                continue
            created = to_turkey_time(w["created"]) if w.get("created") else None
            rows.append(
                {
                    "user": (w.get("author") or {}).get("displayName", "?"),
                    "date": started.date(),
                    "time": started.time().replace(second=0),
                    "key": key,
                    "summary": fields.get("summary", ""),
                    "project": (fields.get("project") or {}).get("key", ""),
                    "type": (fields.get("issuetype") or {}).get("name", ""),
                    "hours": round(w.get("timeSpentSeconds", 0) / 3600, 2),
                    "comment": str(w.get("comment") or "").replace("\r", " ").replace("\n", " ").strip(),
                    "created": created,
                    "url": f"{client.base_url}/browse/{key}",
                }
            )
        say("issues", n, total)

    rows.sort(key=lambda r: (r["user"].lower(), r["date"], r["time"]))
    return rows


# --------------------------------------------------------------------------
# Çıktı dosyaları
# --------------------------------------------------------------------------
def default_filename(projects: list[str], start: date, end: date, ext: str) -> str:
    proj = "-".join(projects) if len(projects) <= 3 else f"{len(projects)}-proje"
    return f"Jira_Worklog_{proj}_{start:%Y-%m-%d}_{end:%Y-%m-%d}.{ext}"


def unique_path(directory: Path, filename: str) -> Path:
    """Aynı isimli dosya varsa (ya da Excel'de açıksa) üstüne yazmamak için ' (2)' ekler."""
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / filename
    stem, suffix, n = path.stem, path.suffix, 2
    while path.exists():
        path = directory / f"{stem} ({n}){suffix}"
        n += 1
    return path


def write_csv(rows: list[dict], path: Path, delimiter: str = ";", decimal_comma: bool = True) -> None:
    with open(path, "w", newline="", encoding="utf-8-sig") as f:  # utf-8-sig: Excel Türkçe karakterleri doğru açar
        w = csv.writer(f, delimiter=delimiter)
        w.writerow(COLUMNS)
        for r in rows:
            hours = str(r["hours"]).replace(".", ",") if decimal_comma else str(r["hours"])
            w.writerow(
                [r["user"], f"{r['date']:%Y-%m-%d}", f"{r['time']:%H:%M}", r["key"], r["summary"],
                 r["project"], r["type"], hours, r["comment"],
                 f"{r['created']:%Y-%m-%d %H:%M}" if r.get("created") else ""]
            )


def write_xlsx(rows: list[dict], path: Path, meta: dict) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wb = Workbook()
    ws = wb.active
    ws.title = "Worklog"
    ws.append(COLUMNS)
    for r in rows:
        ws.append([r["user"], datetime.combine(r["date"], datetime.min.time()), r["time"], r["key"], r["summary"], r["project"], r["type"], r["hours"],
                   r["comment"], r.get("created")])

    last = len(rows) + 1
    for row in ws.iter_rows(min_row=2, max_row=last):
        # Küçük harfli biçim kodları Excel dışı programlarda (Numbers, Önizleme vb.) da tarih olarak tanınır.
        row[1].number_format = "dd.mm.yyyy"
        row[2].number_format = "hh:mm"
        row[7].number_format = "0.00"
        row[9].number_format = "dd.mm.yyyy hh:mm"
        for c in (row[1], row[2], row[9]):
            c.alignment = Alignment(horizontal="left")
        url = rows[row[0].row - 2]["url"]
        row[3].hyperlink = url
        row[3].style = "Hyperlink"

    table = Table(displayName="Worklog", ref=f"A1:{chr(64 + len(COLUMNS))}{last}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"
    for col, width in zip("ABCDEFGHIJ", (24, 12, 14, 14, 52, 10, 14, 8, 60, 18)):
        ws.column_dimensions[col].width = width

    # Özet sayfası: kişi bazında toplam
    summ = wb.create_sheet("Özet")
    summ.append(["Kullanıcı", "Worklog sayısı", "Toplam saat"])
    per_user: dict[str, list] = {}
    for r in rows:
        agg = per_user.setdefault(r["user"], [0, 0.0])
        agg[0] += 1
        agg[1] += r["hours"]
    for user, (count, hours) in sorted(per_user.items(), key=lambda kv: -kv[1][1]):
        summ.append([user, count, round(hours, 2)])
    summ.append(["Toplam", len(rows), round(sum(r["hours"] for r in rows), 2)])
    for cell in summ[1] + summ[summ.max_row]:
        cell.font = Font(bold=True)
    for row in summ.iter_rows(min_row=2):
        row[2].number_format = "0.00"
    summ.column_dimensions["A"].width = 28
    summ.column_dimensions["B"].width = 16
    summ.column_dimensions["C"].width = 14
    info_row = summ.max_row + 2
    for label, value in (
        ("Projeler", ", ".join(meta.get("projects", []))),
        ("Dönem", f"{meta['start']:%d.%m.%Y} – {meta['end']:%d.%m.%Y}"),
        ("Oluşturulma", f"{datetime.now():%d.%m.%Y %H:%M}"),
        ("Jira", meta.get("url", "")),
        ("Saatler", "Türkiye saati (UTC+3)"),
        ("Başlama saati", "Jira'daki «Date Started» alanı: worklog girilirken seçilen başlangıç zamanı"),
        ("Girilme zamanı", "Worklog'un Jira'ya kaydedildiği an"),
    ):
        summ.cell(row=info_row, column=1, value=label).font = Font(bold=True)
        summ.cell(row=info_row, column=2, value=value).alignment = Alignment(horizontal="left")
        info_row += 1

    wb.save(path)
    _mark_number_formats(path)


def _mark_number_formats(path: Path) -> None:
    """Biçimli hücre stillerine applyNumberFormat="1" ekler (openpyxl yazmaz).

    Excel bu bayrağı gözetmeden biçimi uygular; bazı başka programlar (Numbers, Önizleme vb.) ise bayrak yoksa
    biçimi yok sayıp tarihi 46289 gibi ham sayı olarak gösterebilir.
    """
    import re
    import zipfile

    tmp = Path(str(path) + ".tmp")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/styles.xml":
                text = data.decode("utf-8")
                head, sep, tail = text.partition("<cellXfs")
                if sep:
                    body, sep2, rest = tail.partition("</cellXfs>")
                    body = re.sub(r'<xf numFmtId="(?!0")(\d+)"(?![^>]*applyNumberFormat)', r'<xf numFmtId="\1" applyNumberFormat="1"', body)
                    text = head + sep + body + sep2 + rest
                data = text.encode("utf-8")
            zout.writestr(item, data)
    tmp.replace(path)
